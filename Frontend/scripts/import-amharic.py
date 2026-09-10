"""
Run this after you fill in the "Amharic (fill this in)" column in the
translation Excel file and re-upload it.

Usage:
    python3 scripts/import-amharic.py path/to/filled_translations.xlsx

It reads columns: Key | English (source) | Amharic (fill this in) | Used in
and writes src/locales/am.json, keeping any key whose Amharic cell is still
blank as an empty string (so it falls back to English automatically).
"""
import sys
import json
from openpyxl import load_workbook

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 import-amharic.py <translated.xlsx>")
        sys.exit(1)

    path = sys.argv[1]
    wb = load_workbook(path)
    ws = wb.active

    am = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, english, amharic = row[0], row[1], row[2]
        if not key:
            continue
        am[key] = (amharic or "").strip()

    out_path = "src/locales/am.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(am, f, ensure_ascii=False, indent=2)

    filled = sum(1 for v in am.values() if v)
    print(f"Wrote {out_path}: {filled}/{len(am)} keys translated.")

if __name__ == "__main__":
    main()
